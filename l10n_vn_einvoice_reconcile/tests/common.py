# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import io

from openpyxl import Workbook

from odoo.addons.l10n_vn_einvoice.tests.common import COMPANY_VAT, EInvoiceCase

PERIOD = "Từ ngày 01/10/2026 đến ngày 31/10/2026"

# The official layouts (design §4): (key, header) per column, from column A.
LAYOUTS = {
    "ordinary": {
        "title": "DANH SÁCH HÓA ĐƠN",
        "columns": [
            ("stt", "STT"),
            ("form", "Ký hiệu mẫu số"),
            ("symbol", "Ký hiệu hóa đơn"),
            ("number", "Số hóa đơn"),
            ("date", "Ngày lập"),
            ("seller_vat", "MST người bán/MST người xuất hàng"),
            ("seller_name", "Tên người bán/Tên người xuất hàng"),
            ("buyer_vat", "MST người mua/MST người nhận hàng"),
            ("buyer_name", "Tên người mua/Tên người nhận hàng"),
            ("buyer_address", "Địa chỉ người mua"),
            ("untaxed", "Tổng tiền chưa thuế"),
            ("tax", "Tổng tiền thuế"),
            ("discount", "Tổng tiền chiết khấu thương mại"),
            ("fees", "Tổng tiền phí"),
            ("total", "Tổng tiền thanh toán"),
            ("currency", "Đơn vị tiền tệ"),
            ("rate", "Tỷ giá"),
            ("status", "Trạng thái hóa đơn"),
            ("check", "Kết quả kiểm tra hóa đơn"),
        ],
    },
    "cash_register": {
        "title": "DANH SÁCH HÓA ĐƠN MÁY TÍNH TIỀN",
        "columns": [
            ("stt", "STT"),
            ("form", "Ký hiệu mẫu số"),
            ("symbol", "Ký hiệu hóa đơn"),
            ("number", "Số hóa đơn"),
            ("date", "Ngày lập"),
            ("seller_vat", "MST người bán/MST người xuất hàng"),
            ("seller_name", "Tên người bán/Tên người xuất hàng"),
            ("seller_address", "Địa chỉ người bán"),
            ("buyer_vat", "MST người mua/MST người nhận hàng"),
            ("buyer_name", "Tên người mua/Tên người nhận hàng"),
            ("citizen_id", "Căn cước công dân"),
            ("untaxed", "Tổng tiền chưa thuế"),
            ("tax", "Tổng tiền thuế"),
            ("discount", "Tổng tiền chiết khấu thương mại"),
            ("total", "Tổng tiền thanh toán"),
            ("status", "Trạng thái hóa đơn"),
            ("check", "Kết quả kiểm tra hóa đơn"),
        ],
    },
}


class PortalListCase(EInvoiceCase):
    """Made-up portal lists in the official layouts; no test methods."""

    def _row(self, **values):
        row = {
            "stt": 1.0,
            "form": "1",
            "symbol": "C26TAA",
            "number": "100",
            "date": "01/10/2026",
            "seller_vat": "0314071202",
            "seller_name": "Nhà cung cấp Thử",
            "seller_address": "1 Đường Thử",
            "buyer_vat": COMPANY_VAT,
            "buyer_name": "CÔNG TY TNHH DV TM QUANG PHƯƠNG",
            "buyer_address": "244 Hồ Học Lãm",
            "citizen_id": None,
            "untaxed": 1000000.0,
            "tax": 80000.0,
            "discount": 0.0,
            "fees": None,
            "total": 1080000.0,
            "currency": "VND",
            "rate": "1.0",
            "status": "Hóa đơn mới",
            "check": "Đã cấp mã hóa đơn",
        }
        row.update(values)
        return row

    def _portal_list(self, rows, kind="ordinary", title=None, period=PERIOD, headers=None):
        """An .xlsx in the kind's layout, base64-encoded: the title in A3, the
        period in A4, the headers on row 6, the rows from row 7."""
        layout = LAYOUTS[kind]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet 1"
        sheet["A3"] = layout["title"] if title is None else title
        sheet["A4"] = period
        for index, (key, header) in enumerate(layout["columns"], start=1):
            sheet.cell(row=6, column=index, value=(headers or {}).get(key, header))
        for offset, row in enumerate(rows):
            for index, (key, _header) in enumerate(layout["columns"], start=1):
                sheet.cell(row=7 + offset, column=index, value=row.get(key))
        stream = io.BytesIO()
        workbook.save(stream)
        return base64.b64encode(stream.getvalue())

    def _new_batch(self, direction="purchase", kind="ordinary", user=None):
        env = self.env(user=user) if user else self.env
        return env["l10n_vn.einvoice.batch"].create(
            {"direction": direction, "list_kind": kind}
        )

    def _import(self, batch, content, filename="list.xlsx"):
        wizard = self.env["l10n_vn.einvoice.batch.import"].with_env(batch.env).create(
            {"batch_id": batch.id, "file": content, "filename": filename}
        )
        wizard.action_import()
        batch.invalidate_recordset()
        return batch

    def _batch(self, direction, content, kind="ordinary"):
        return self._import(self._new_batch(direction, kind), content)
